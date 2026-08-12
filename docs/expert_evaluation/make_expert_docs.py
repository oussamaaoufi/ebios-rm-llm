"""
make_expert_docs.py
-------------------
Generates one Word (.docx) evaluation booklet per evaluator from
expert_package_blind.json. Each scenario's JSON is rendered as clean,
labelled, readable content (NOT raw JSON), followed by a scoring table.

Requires: pip install python-docx

Run:
    python scripts/evaluation\make_expert_docs.py
"""

import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from config import BASE_DIR as BASE
EVAL = BASE / "evaluation_results"
OUT = EVAL / "expert_docs"
OUT.mkdir(exist_ok=True)

EVALUATORS = ["E1", "E2", "E3"]
NAVY = RGBColor(0x1F, 0x38, 0x64)

CRITERIA = [
    ("Technical accuracy", "MITRE ATT&CK techniques correct and appropriate"),
    ("EBIOS-RM compliance", "Follows the canonical structure for the workshop"),
    ("Contextual relevance", "Genuinely tailored to the stated sector"),
    ("Attack path realism", "Attack chain is plausible and coherent"),
    ("Completeness", "All required elements present with actionable detail"),
]

blind = json.loads((EVAL / "expert_package_blind.json").read_text(encoding="utf-8"))


def label(key):
    return str(key).replace("_", " ").replace("-", " ").strip().capitalize()


def add_scenario(doc, obj, level=0):
    """Recursively render a parsed JSON scenario as readable labelled content."""
    indent = Inches(0.25 * level)
    if isinstance(obj, dict):
        for k, v in obj.items():
            if isinstance(v, (dict, list)):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = indent
                p.paragraph_format.space_before = Pt(4)
                run = p.add_run(label(k))
                run.bold = True
                run.font.size = Pt(10.5)
                run.font.color.rgb = NAVY
                add_scenario(doc, v, level + 1)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = indent
                p.paragraph_format.space_after = Pt(1)
                r1 = p.add_run(f"{label(k)}: ")
                r1.bold = True
                r1.font.size = Pt(10)
                r2 = p.add_run(str(v))
                r2.font.size = Pt(10)
    elif isinstance(obj, list):
        for i, item in enumerate(obj, 1):
            if isinstance(item, dict):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = indent
                run = p.add_run(f"• Item {i}")
                run.bold = True
                run.italic = True
                run.font.size = Pt(9.5)
                add_scenario(doc, item, level + 1)
            else:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = indent
                p.add_run(str(item)).font.size = Pt(10)


def scoring_table(doc):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "Criterion"; hdr[1].text = "Score (1–5)"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True; r.font.size = Pt(10)
    t.columns[0].width = Inches(4.5)
    t.columns[1].width = Inches(1.3)
    for name, desc in CRITERIA:
        row = t.add_row().cells
        p = row[0].paragraphs[0]
        r1 = p.add_run(name); r1.bold = True; r1.font.size = Pt(10)
        r2 = p.add_run(f"  — {desc}"); r2.italic = True; r2.font.size = Pt(8.5)
        row[1].text = ""
    total = t.add_row().cells
    rp = total[0].paragraphs[0].add_run("TOTAL  (out of 25)")
    rp.bold = True; rp.font.size = Pt(10)
    total[1].text = ""
    return t


def build(code):
    doc = Document()
    # base style
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # ---- cover / instructions ----
    h = doc.add_paragraph()
    r = h.add_run("EBIOS-RM Scenario Evaluation")
    r.bold = True; r.font.size = Pt(18); r.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    rs = sub.add_run(f"Evaluator: {code}   ·   20 scenarios   ·   approx. 90 minutes")
    rs.italic = True; rs.font.size = Pt(11)

    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.add_run(
        "Please score each of the 20 scenarios below on five criteria, from 1 (poor) to "
        "5 (excellent), by writing a number in the score column of each table. You may add a "
        "brief comment beneath each. The scenarios come from several sources in random order; "
        "please score only what is in front of you. Use the full 1–5 range."
    ).font.size = Pt(10.5)

    # rubric summary
    doc.add_paragraph()
    rh = doc.add_paragraph().add_run("Scoring guide")
    rh.bold = True; rh.font.size = Pt(12); rh.font.color.rgb = NAVY
    guide = doc.add_paragraph()
    guide.add_run("1 = Poor    3 = Acceptable    5 = Excellent    (2 and 4 for in-between).").font.size = Pt(10)
    for name, desc in CRITERIA:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{name}: "); r1.bold = True; r1.font.size = Pt(10)
        p.add_run(desc).font.size = Pt(10)

    doc.add_page_break()

    # ---- scenarios ----
    for idx, item in enumerate(blind, 1):
        title = doc.add_paragraph()
        tr = title.add_run(f"Scenario {item['item_id']}")
        tr.bold = True; tr.font.size = Pt(14); tr.font.color.rgb = NAVY
        title.paragraph_format.space_before = Pt(6)

        # instruction (the task the scenario answers)
        ip = doc.add_paragraph()
        ir = ip.add_run("Task:  ")
        ir.bold = True; ir.font.size = Pt(10)
        ip.add_run(item["instruction"]).font.size = Pt(10)
        ip.paragraph_format.space_after = Pt(6)

        # scenario content, rendered readable
        sh = doc.add_paragraph().add_run("Scenario")
        sh.bold = True; sh.font.size = Pt(11); sh.font.color.rgb = NAVY
        scen = item["scenario"].strip()
        parsed = None
        if "{" in scen and "}" in scen:
            try:
                parsed = json.loads(scen[scen.find("{"): scen.rfind("}") + 1])
            except Exception:
                parsed = None
        if parsed is not None:
            add_scenario(doc, parsed)
        else:
            doc.add_paragraph(scen).runs[0].font.size = Pt(10)

        # scoring table
        doc.add_paragraph()
        sc = doc.add_paragraph().add_run("Your scores")
        sc.bold = True; sc.font.size = Pt(10.5); sc.font.color.rgb = NAVY
        scoring_table(doc)
        cp = doc.add_paragraph()
        cr = cp.add_run("Comment / what would you change?  ")
        cr.bold = True; cr.font.size = Pt(9.5)
        cp.add_run("_______________________________________________").font.size = Pt(9.5)

        if idx < len(blind):
            doc.add_page_break()

    dest = OUT / f"EBIOS_Expert_Evaluation_{code}.docx"
    doc.save(dest)
    return dest


for code in EVALUATORS:
    d = build(code)
    print(f"  wrote {d.name}")

print(f"\n{len(EVALUATORS)} Word booklets in {OUT}")
print("Scenarios rendered as readable labelled content with a scoring table each.")
print("Send one per evaluator. Do NOT send expert_package_key.json.")
